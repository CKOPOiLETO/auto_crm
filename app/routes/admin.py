# app/routes/admin.py
from flask import Blueprint, render_template, request, flash, redirect, url_for
from app import db
from app.models.tariff import Tariff
from flask_login import login_required, current_user
from functools import wraps
from flask import abort
from app.models.user import User
from werkzeug.security import generate_password_hash
from sqlalchemy import or_
from sqlalchemy import func, desc # Для подсчета, суммирования и сортировки
from app.models.proposal import Proposal
from app.models.car import Car
from sqlalchemy import func, desc, case
from app.models.client import Client
from app.models.proposal import Proposal
from app.models.car import Car
from datetime import datetime, time, timedelta
from datetime import datetime, timedelta
from sqlalchemy import func, desc, case, cast, Date
import csv
import io
from flask import Response
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import io

def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not current_user.is_authenticated or current_user.role != 'admin':
            abort(403)
        return f(*args, **kwargs)
    return decorated_function

# Обязательно импортируем нашу функцию парсера валют
from app.services.currency_updater import get_nbrb_rates

admin_bp = Blueprint('admin', __name__, url_prefix='/admin', template_folder='../templates/admin')
@admin_bp.route('/tariffs', methods=['GET', 'POST'])
@login_required
@admin_required
def tariffs():
    tariff = Tariff.query.first()
    
    if not tariff:
        # Стартовые значения
        tariff = Tariff(
            usd_rate=3.2500, 
            eur_rate=3.5500, 
            shipping_usa=600.00, 
            shipping_sea=1400.00,
            shipping_eu=500.00, 
            auction_fee_rate=5.00
        )
        db.session.add(tariff)
        db.session.commit()

    # Кусок файла app/routes/admin.py
    if request.method == 'POST':
        try:
            tariff.usd_rate = request.form.get('usd_rate')
            tariff.eur_rate = request.form.get('eur_rate')
            tariff.shipping_usa = request.form.get('shipping_usa')
            tariff.shipping_sea = request.form.get('shipping_sea')
            tariff.shipping_eu = request.form.get('shipping_eu')
            tariff.auction_fee_rate = request.form.get('auction_fee_rate')
            
            db.session.commit()
            flash('Тарифы успешно обновлены!', 'success')
            return redirect(url_for('admin.tariffs'))
        except Exception as e:
            # --- ИСПРАВЛЕНИЕ: ОТМЕНЯЕМ ОШИБОЧНУЮ ТРАНЗАКЦИЮ ---
            db.session.rollback() 
            flash(f'Ошибка при сохранении: {str(e)}', 'danger')

    return render_template('tariffs.html', tariff=tariff)

@admin_bp.route('/tariffs/update-nbrb', methods=['POST'])
@login_required
@admin_required
def update_nbrb():
    try:
        tariff = Tariff.query.first()
        if not tariff:
            flash('Сначала сохраните базовые тарифы!', 'warning')
            return redirect(url_for('admin.tariffs'))
            
        rates = get_nbrb_rates()
        
        if rates:
            tariff.usd_rate = rates['usd']
            tariff.eur_rate = rates['eur']
            db.session.commit()
            
            if rates['date'] != 'Error':
                flash(f"Курсы на {rates['date']} успешно загружены! USD: {rates['usd']}, EUR: {rates['eur']}", 'success')
            else:
                flash(f"Сайт НБРБ недоступен. Установлены резервные курсы: USD {rates['usd']}, EUR {rates['eur']}", 'warning')
        else:
            flash('Не удалось получить данные. Проверьте консоль.', 'danger')
            
    except Exception as e:
        # Если произойдет сбой в БД или коде, мы не получим 500 ошибку, а увидим красную плашку
        flash(f'Внутренняя ошибка сервера: {str(e)}', 'danger')

    return redirect(url_for('admin.tariffs'))


# ==========================================
# УПРАВЛЕНИЕ СОТРУДНИКАМИ
# ==========================================

@admin_bp.route('/users', methods=['GET', 'POST'])
@login_required
@admin_required
def manage_users():
    if request.method == 'POST':
        login_input = request.form.get('login')
        password_input = request.form.get('password')
        full_name_input = request.form.get('full_name')
        
        existing_user = User.query.filter_by(login=login_input).first()
        
        if existing_user:
            flash(f'Ошибка: Сотрудник с логином "{login_input}" уже существует!', 'danger')
        else:
            try:
                new_user = User(
                    login=login_input,
                    password_hash=generate_password_hash(password_input),
                    full_name=full_name_input,
                    role='manager' # Жестко задаем роль менеджера
                )
                db.session.add(new_user)
                db.session.commit()
                flash(f'Сотрудник {full_name_input} успешно добавлен!', 'success')
            except Exception as e:
                db.session.rollback()
                flash(f'Произошла ошибка при сохранении: {str(e)}', 'danger')
                
        return redirect(url_for('admin.manage_users'))

    q = request.args.get('q', '')
    query = User.query
    if q:
        term = f"%{q}%"
        from sqlalchemy import or_
        query = query.filter(or_(User.login.ilike(term), User.full_name.ilike(term)))
    
    users = query.order_by(User.full_name).all()
    return render_template('admin/users.html', users=users, q=q)


@admin_bp.route('/users/edit/<int:user_id>', methods=['GET', 'POST'])
@login_required
@admin_required
def edit_user(user_id):
    user = User.query.get_or_404(user_id)
    
    if request.method == 'POST':
        user.login = request.form.get('login')
        user.full_name = request.form.get('full_name')
        
        # Обновляем пароль, только если он был введен
        new_password = request.form.get('password')
        if new_password:
            user.password_hash = generate_password_hash(new_password)
            
        db.session.commit()
        flash('Данные сотрудника успешно обновлены!', 'info')
        return redirect(url_for('admin.manage_users'))
        
    return render_template('admin/user_edit.html', user=user)


@admin_bp.route('/users/delete/<int:user_id>', methods=['POST'])
@login_required
@admin_required
def delete_user(user_id):
    if current_user.id == user_id:
        flash('Вы не можете удалить свою собственную учетную запись!', 'danger')
        return redirect(url_for('admin.manage_users'))
        
    user = User.query.get_or_404(user_id)
    db.session.delete(user)
    db.session.commit()
    flash(f'Сотрудник {user.login} удален из системы.', 'warning')
    return redirect(url_for('admin.manage_users'))





@admin_bp.route('/analytics')
@login_required
@admin_required
def analytics():
    # ========================================================
    # 1. Заявки в разрезе статуса (С ФИЛЬТРОМ ПО ДАТАМ)
    # ========================================================
    status_start_str = request.args.get('status_date_from')
    status_end_str = request.args.get('status_date_to')
    
    status_query = db.session.query(Proposal.status, func.count(Proposal.id))
    
    # Применяем фильтры дат, если админ их выбрал
    if status_start_str:
        st_date = datetime.strptime(status_start_str, '%Y-%m-%d').date()
        status_query = status_query.filter(cast(Proposal.created_at, Date) >= st_date)
    if status_end_str:
        en_date = datetime.strptime(status_end_str, '%Y-%m-%d').date()
        status_query = status_query.filter(cast(Proposal.created_at, Date) <= en_date)
        
    status_counts = status_query.group_by(Proposal.status).all()
    status_labels = [s[0] for s in status_counts]
    status_data = [s[1] for s in status_counts]


    # ========================================================
    # 2. Конверсия во времени (С ФИЛЬТРОМ ПО ДАТАМ)
    # ========================================================
    start_str = request.args.get('start_date')
    end_str = request.args.get('end_date')
    today = datetime.now().date()
    
    try:
        end_date = datetime.strptime(end_str, '%Y-%m-%d').date() if end_str else today
        start_date = datetime.strptime(start_str, '%Y-%m-%d').date() if start_str else (end_date - timedelta(days=13))
    except ValueError:
        end_date = today
        start_date = end_date - timedelta(days=13)

    if start_date > end_date:
        start_date, end_date = end_date, start_date

    daily_stats = db.session.query(
        cast(Proposal.created_at, Date).label('date'),
        func.sum(case((Proposal.status.in_(['sent', 'accepted', 'rejected']), 1), else_=0)).label('sent_count'),
        func.sum(case((Proposal.status == 'accepted', 1), else_=0)).label('accepted_count')
    ).filter(
        cast(Proposal.created_at, Date) >= start_date,
        cast(Proposal.created_at, Date) <= end_date
    ).group_by(cast(Proposal.created_at, Date)).order_by(cast(Proposal.created_at, Date)).all()

    stats_dict = {stat.date: {'sent': stat.sent_count or 0, 'accepted': stat.accepted_count or 0} for stat in daily_stats}

    chart_dates = []
    chart_sent = []
    chart_accepted = []

    # Рисуем график строго по выбранному диапазону
    delta_days = (end_date - start_date).days
    for i in range(delta_days + 1):
        current_d = start_date + timedelta(days=i)
        chart_dates.append(current_d.strftime('%d.%m'))
        chart_sent.append(int(stats_dict.get(current_d, {}).get('sent', 0)))
        chart_accepted.append(int(stats_dict.get(current_d, {}).get('accepted', 0)))

    total_sent = sum(chart_sent)
    total_acc = sum(chart_accepted)
    conversion_rate = round((total_acc / total_sent * 100), 2) if total_sent > 0 else 0


    # ========================================================
    # 3. Эффективность менеджеров и Выручка
    # ========================================================
    manager_stats = db.session.query(
        User.full_name,
        func.count(Proposal.id).label('total_proposals'),
        func.sum(case((Proposal.status == 'sent', 1), else_=0)).label('pending_proposals'),
        func.sum(case((Proposal.status == 'accepted', 1), else_=0)).label('accepted_proposals'),
        func.sum(case((Proposal.status == 'accepted', Proposal.total_price_usd), else_=0)).label('total_revenue')
    ).join(Client, Client.manager_id == User.id)\
     .join(Proposal, Proposal.client_id == Client.id)\
     .group_by(User.id, User.full_name).order_by(desc('total_revenue')).all()

    total_revenue = db.session.query(func.sum(Proposal.total_price_usd)).filter_by(status='accepted').scalar() or 0

    return render_template('admin/analytics.html', 
                           conversion_rate=conversion_rate,
                           total_revenue=total_revenue,
                           status_labels=status_labels,
                           status_data=status_data,
                           manager_stats=manager_stats,
                           chart_dates=chart_dates,
                           chart_sent=chart_sent,
                           chart_accepted=chart_accepted,
                           # Возвращаем даты обратно в HTML
                           start_date=start_date.strftime('%Y-%m-%d'),
                           end_date=end_date.strftime('%Y-%m-%d'),
                           status_date_from=status_start_str or '',
                           status_date_to=status_end_str or ''
                           )


@admin_bp.route('/hierarchy')
@login_required
@admin_required
def hierarchy():
    """Древовидная структура: Менеджеры -> Их Клиенты -> Их Предложения"""
    
    # Вытягиваем всех пользователей вместе со всеми их клиентами, предложениями и авто
    users = User.query.options(
        db.joinedload(User.clients)
          .joinedload(Client.proposals)
          .joinedload(Proposal.car)
    ).order_by(User.full_name).all()
    
    return render_template('admin/hierarchy.html', users=users)





# --- ЭКСПОРТ ОТЧЕТА В НАСТОЯЩИЙ EXCEL (.XLSX) ---
@admin_bp.route('/analytics/export_csv')
@login_required
@admin_required
def export_analytics_csv():
    # 1. Запрашиваем данные
    manager_stats = db.session.query(
        User.full_name,
        func.count(Proposal.id).label('total_proposals'),
        func.sum(case((Proposal.status == 'sent', 1), else_=0)).label('pending_proposals'),
        func.sum(case((Proposal.status == 'accepted', 1), else_=0)).label('accepted_proposals'),
        func.sum(case((Proposal.status == 'accepted', Proposal.total_price_usd), else_=0)).label('total_revenue')
    ).join(Client, Client.manager_id == User.id)\
     .join(Proposal, Proposal.client_id == Client.id)\
     .group_by(User.id, User.full_name).order_by(desc('total_revenue')).all()

    # 2. Создаем Excel-книгу и активный лист
    wb = Workbook()
    ws = wb.active
    ws.title = "Эффективность менеджеров"

    # 3. Пишем заголовки и делаем их ЖИРНЫМИ и по центру
    headers = ['Сотрудник', 'Создано КП', 'Ожидают ответа', 'Сделок закрыто', 'Конверсия', 'Сумма сделок (USD)']
    ws.append(headers)
    
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')

    # 4. Записываем данные (теперь как настоящие числа, а не текст!)
    for m in manager_stats:
        total = m.total_proposals or 0
        accepted = m.accepted_proposals or 0
        revenue = float(m.total_revenue or 0)
        conversion = (accepted / total) if total > 0 else 0 # Просто число, Excel сам сделает проценты

        ws.append([
            m.full_name,
            total,
            m.pending_proposals or 0,
            accepted,
            conversion,
            revenue
        ])

    # 5. Красивое форматирование цифр и ширины колонок
    for row_idx in range(2, len(manager_stats) + 2):
        ws[f"E{row_idx}"].number_format = '0.0%' # Ячейка Конверсии: формат процентов
        ws[f"F{row_idx}"].number_format = '#,##0.00' # Ячейка Суммы: разделитель тысяч

    # Задаем фиксированную ширину колонок, чтобы всё помещалось
    ws.column_dimensions['A'].width = 35 # Сотрудник
    ws.column_dimensions['B'].width = 15 # Создано
    ws.column_dimensions['C'].width = 18 # Ожидают
    ws.column_dimensions['D'].width = 18 # Сделок
    ws.column_dimensions['E'].width = 15 # Конверсия
    ws.column_dimensions['F'].width = 25 # Сумма

    # 6. Сохраняем в бинарный поток
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    # 7. Отдаем пользователю файл формата .xlsx
    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Managers_Efficiency_Report.xlsx"}
    )

# --- ЭКСПОРТ ДАННЫХ ДИАГРАММЫ СТАТУСОВ ---
@admin_bp.route('/analytics/export_status_csv')
@login_required
@admin_required
def export_status_csv():
    status_counts = db.session.query(Proposal.status, func.count(Proposal.id)).group_by(Proposal.status).all()
    
    output = io.StringIO()
    output.write('\ufeff')
    writer = csv.writer(output, delimiter=';')
    writer.writerow(['Статус предложения', 'Количество (шт.)'])
    
    # Словарь для красивого перевода статусов в файле Excel
    status_map = {'draft': 'Черновик', 'sent': 'Отправлено', 'accepted': 'Принято', 'rejected': 'Отказ'}
    
    for s in status_counts:
        status_ru = status_map.get(s[0], s[0])
        writer.writerow([status_ru, s[1]])
        
    return Response(output.getvalue(), mimetype="text/csv", headers={"Content-Disposition": "attachment; filename=Status_Report.csv"})


# --- ЭКСПОРТ ДАННЫХ КОНВЕРСИИ ВО ВРЕМЕНИ ---
@admin_bp.route('/analytics/export_conversion_csv')
@login_required
@admin_required
def export_conversion_csv():
    # Берем даты из параметров URL (чтобы выгрузить именно выбранный период)
    start_str = request.args.get('start_date')
    end_str = request.args.get('end_date')
    today = datetime.now().date()
    
    try:
        end_date = datetime.strptime(end_str, '%Y-%m-%d').date() if end_str else today
        start_date = datetime.strptime(start_str, '%Y-%m-%d').date() if start_str else (end_date - timedelta(days=13))
    except ValueError:
        end_date = today
        start_date = end_date - timedelta(days=13)

    if start_date > end_date:
        start_date, end_date = end_date, start_date

    daily_stats = db.session.query(
        cast(Proposal.created_at, Date).label('date'),
        func.sum(case((Proposal.status.in_(['sent', 'accepted', 'rejected']), 1), else_=0)).label('sent_count'),
        func.sum(case((Proposal.status == 'accepted', 1), else_=0)).label('accepted_count')
    ).filter(
        cast(Proposal.created_at, Date) >= start_date,
        cast(Proposal.created_at, Date) <= end_date
    ).group_by(cast(Proposal.created_at, Date)).order_by(cast(Proposal.created_at, Date)).all()

    stats_dict = {stat.date: {'sent': stat.sent_count or 0, 'accepted': stat.accepted_count or 0} for stat in daily_stats}

    output = io.StringIO()
    output.write('\ufeff')
    writer = csv.writer(output, delimiter=';')
    writer.writerow(['Дата', 'Отправлено предложений', 'Успешных сделок (Принято)'])

    delta_days = (end_date - start_date).days
    for i in range(delta_days + 1):
        current_d = start_date + timedelta(days=i)
        sent = int(stats_dict.get(current_d, {}).get('sent', 0))
        accepted = int(stats_dict.get(current_d, {}).get('accepted', 0))
        writer.writerow([current_d.strftime('%d.%m.%Y'), sent, accepted])

    return Response(output.getvalue(), mimetype="text/csv", headers={"Content-Disposition": "attachment; filename=Conversion_Timeline_Report.csv"})